# -*- coding: utf-8 -*-
"""
One reading of RMTH_Forms_and_Calculations_v2.xlsx -- the workbook of
1 October 2026 that replaced Ramtha's seventeen indicator forms (v1/,
0122-0133) with seven operational forms.

Every generator under supabase/ramtha/ imports this and nothing else reads
the workbook, so the option lists, the form definitions and both locale files
come from the SAME cells.

The workbook has two sheets:

  Forms Needed        one row per field: Form ID | sub-page En | sub-page Ar |
                      Field ID | Field Label En | Field Label Ar | Field Type |
                      read-only | Required | Options En ("-" delimiter) |
                      Options Ar ("-" delimiter) | Validation | Help Text |
                      Dependency | Page En | Page Ar | Indicator ID |
                      Mapping role
  Calculation Method  one row per indicator: Indicator ID | Indicator |
                      Required Field(s) | Calculation Type | Formula |
                      Output Unit

Every string is kept VERBATIM (whitespace-normalised only); nothing here
translates, drafts or corrects. What the platform adds beyond the sheet --
the two fields the owner asked for, the Arabic of the help texts -- is
catalogue.py's, each one named there.
"""
import os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..'))
XLSX = os.path.join(ROOT, 'RMTH_Forms_and_Calculations_v2.xlsx')


def norm(s):
    if s is None:
        return ''
    return re.sub(r'[ \t\xa0]+', ' ', str(s)).strip()


def split_options(s):
    """An options cell: the sheet's own delimiter is ' - '."""
    s = norm(s)
    return [norm(x) for x in s.split(' - ')] if s else []


# A full indicator code as the sheet writes it: RMTH-IMP-0, RMTH-SO1-0,
# RMTH-SO1-A1, RMTH-SO1-A0.1, RMTH-SO2-C1.2 ...
CODE_RE = re.compile(r'RMTH-(?:IMP-0|SO\d-(?:0|[A-Z]\d+(?:\.\d+)?))')


class Field:
    __slots__ = ('form', 'fid', 'sub_en', 'sub_ar', 'label_en', 'label_ar', 'ftype', 'readonly',
                 'required', 'opts_en', 'opts_ar', 'validation', 'help', 'dependency', 'page_en',
                 'page_ar', 'indicators', 'indicator_text', 'role', 'row')

    def __repr__(self):
        return 'Field(%s %s %s)' % (self.form, self.fid, self.ftype)


class Form:
    __slots__ = ('form', 'fields')

    def __init__(self, form):
        self.form = form
        self.fields = []

    def field(self, fid):
        for f in self.fields:
            if f.fid == fid:
                return f
        raise KeyError('%s has no field %s' % (self.form, fid))


def load_forms():
    """{ 'FORM-01': Form, ... } in the sheet's order, fields in the sheet's order."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Forms Needed']
    header = [norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    expect = ['Form ID', 'sub-page En', 'sub-page Ar', 'Field ID', 'Field Label En', 'Field Label Ar',
              'Field Type', 'read-only', 'Required', 'Options En ("-" delimiter)', 'Options Ar ("-" delimiter)',
              'Validation', 'Help Text', 'Dependency', 'Page En', 'Page Ar', 'Indicator ID', 'Mapping role']
    if header[:len(expect)] != expect:
        raise SystemExit('Forms Needed: the header is not the one this reader was written for:\n%r' % header)
    forms = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[0]:
            continue
        r = list(r) + [None] * (18 - len(r))
        f = Field()
        f.row = i
        f.form = norm(r[0])
        f.sub_en, f.sub_ar = norm(r[1]), norm(r[2])
        f.fid = norm(r[3])
        f.label_en, f.label_ar = norm(r[4]), norm(r[5])
        f.ftype = norm(r[6])
        f.readonly = norm(r[7]).lower() == 'yes'
        f.required = norm(r[8]).lower() == 'yes'
        f.opts_en, f.opts_ar = split_options(r[9]), split_options(r[10])
        f.validation, f.help, f.dependency = norm(r[11]), norm(r[12]), norm(r[13])
        f.page_en, f.page_ar = norm(r[14]), norm(r[15])
        f.indicator_text = norm(r[16])
        f.indicators = CODE_RE.findall(f.indicator_text)
        f.role = norm(r[17])
        if len(f.opts_en) != len(f.opts_ar):
            raise SystemExit('%s %s: %d English options and %d Arabic ones' % (f.form, f.fid, len(f.opts_en), len(f.opts_ar)))
        forms.setdefault(f.form, Form(f.form)).fields.append(f)
    return forms


class Calc:
    __slots__ = ('code', 'statement', 'fields', 'ctype', 'formula', 'unit', 'row')


def load_calculations():
    """[Calc, ...] in the sheet's order."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Calculation Method']
    header = [norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    if header[:6] != ['Indicator ID', 'Indicator', 'Required Field(s)', 'Calculation Type', 'Formula', 'Output Unit']:
        raise SystemExit('Calculation Method: unexpected header %r' % header)
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[0]:
            continue
        c = Calc()
        c.row = i
        c.code = norm(r[0])
        c.statement = norm(r[1])
        c.fields = norm(r[2])
        c.ctype = norm(r[3])
        # the formula keeps its line breaks: each line is a clause of the rule
        c.formula = '\n'.join(norm(x) for x in str(r[4] or '').split('\n') if norm(x))
        c.unit = norm(r[5])
        out.append(c)
    return out
