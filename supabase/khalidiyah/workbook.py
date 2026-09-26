# -*- coding: utf-8 -*-
"""
One reading of Khaldia_2_reviewed.xlsx -- the reviewed Khalidiyah workbook of
26 September 2026 that replaced the 21 indicator forms (v1/, 0141-0152) with
24 operational forms.

Every generator under supabase/khalidiyah/ imports this and nothing else reads
the workbook, so the option lists, the tables, the form definitions and both
locale files come from the SAME cells.

The workbook has four sheets:

  Forms needed          one row per field: Form ID | sub-page En | sub-page Ar |
                        Field ID | Field Label En | Field Label Ar | Field Type |
                        read-only | Required | Options En | Options Ar |
                        Validation | Help Text | Dependency | Page En | Page Ar |
                        Indicator ID | Mapping role | Review status |
                        Reviewer note / rationale
  Calculation formulas  one row per indicator: inclusion filter, numerator,
                        denominator, formula, disaggregation, notes -- in
                        terms of the Field IDs above
  Indicator coverage    one row per indicator: which forms and fields feed it
  English_form          the framework (Khaldia.xlsx with one column added)

Every string is kept VERBATIM (whitespace-normalised only); nothing here
translates, drafts or corrects. The corrections the reviewer asks for are the
catalogue's (catalogue.REVIEW_FIXES), each one named, so the sheet's own text
stays readable next to what the app shows.
"""
import os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..'))
XLSX = os.path.join(ROOT, 'Khaldia_2_reviewed.xlsx')


def norm(s):
    if s is None:
        return ''
    return re.sub(r'[ \t\xa0]+', ' ', str(s)).strip()


def lines(s):
    """A multi-line cell as its non-empty lines, each normalised."""
    return [norm(x) for x in str(s or '').split('\n') if norm(x)]


class Field:
    __slots__ = ('form', 'fid', 'sub_en', 'sub_ar', 'label_en', 'label_ar', 'ftype', 'readonly',
                 'required', 'opts_en', 'opts_ar', 'validation', 'help', 'dependency', 'page_en',
                 'page_ar', 'indicators', 'role', 'review', 'note', 'row')

    def __repr__(self):
        return 'Field(%s %s %s)' % (self.form, self.fid, self.ftype)


class Form:
    __slots__ = ('form', 'fields')

    def __init__(self, form):
        self.form = form
        self.fields = []

    def field(self, fid):
        for f in self.fields:
            if f.fid == fid:
                return f
        raise KeyError('%s has no field %s' % (self.form, fid))


def load_forms():
    """{ 'FORM-01': Form, ... } in the sheet's order, fields in the sheet's order."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Forms needed']
    forms = {}
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[0]:
            continue
        r = list(r) + [None] * (20 - len(r))
        f = Field()
        f.row = i
        f.form = norm(r[0])
        f.sub_en, f.sub_ar = norm(r[1]), norm(r[2])
        f.fid = norm(r[3])
        f.label_en, f.label_ar = norm(r[4]), norm(r[5])
        f.ftype = norm(r[6])
        f.readonly = norm(r[7]).lower() == 'yes'
        f.required = norm(r[8]).lower() == 'yes'
        # the option cells keep their raw text: the catalogue types each list
        # out and model.py holds every typed option to these cells
        f.opts_en, f.opts_ar = str(r[9] or ''), str(r[10] or '')
        f.validation, f.help, f.dependency = norm(r[11]), norm(r[12]), norm(r[13])
        f.page_en, f.page_ar = norm(r[14]), norm(r[15])
        f.indicators = [x.strip() for x in norm(r[16]).split(';') if x.strip()]
        f.role, f.review, f.note = norm(r[17]), norm(r[18]), norm(r[19])
        forms.setdefault(f.form, Form(f.form)).fields.append(f)
    return forms


CALC_COLUMNS = ('code', 'indicator', 'type', 'unit', 'frequency', 'source', 'filter',
                'numerator', 'denominator', 'formula', 'disaggregation', 'notes')


def load_calculations():
    """{ 'KHLD-IMP-0': {code, indicator, type, unit, ...} } from 'Calculation formulas'."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Calculation formulas']
    out = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r or not r[0] or not str(r[0]).startswith('KHLD-'):
            continue
        out[norm(r[0])] = dict(zip(CALC_COLUMNS, [norm(x) for x in r[:12]]))
    return out


def load_shared_definitions():
    """The 'Shared definitions' block under the formulas: [(name, text)]."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Calculation formulas']
    out, on = [], False
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == 'Shared definitions':
            on = True
            continue
        if on and r and r[0]:
            out.append((norm(r[0]), norm(r[1])))
    return out


def load_coverage():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['Indicator coverage']
    rows = list(ws.iter_rows(values_only=True))
    head = [norm(x) for x in rows[0]]
    return {norm(r[0]): dict(zip(head, [norm(x) for x in r])) for r in rows[1:] if r and r[0] and str(r[0]).startswith('KHLD-')}


def load_framework():
    """The English_form sheet: header, rows (each a list of normalised cells)."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['English_form']
    rows = [[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    header, body = rows[0], [r for r in rows[1:] if any(r)]
    return header, body


if __name__ == '__main__':
    forms = load_forms()
    n = 0
    for k, f in forms.items():
        n += len(f.fields)
        print(k, len(f.fields), 'fields', f.fields[0].sub_en)
    print('total fields', n, '| indicators with a formula', len(load_calculations()))
